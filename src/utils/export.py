"""Export functionality for solutions and visualizations."""

from fpdf import FPDF
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from typing import Dict, List
import matplotlib.pyplot as plt
from pathlib import Path


class PDFExporter:
    """Export solutions to PDF format."""
    
    def export_solution(self, data: Dict, output_path: str):
        """
        Export a solution to PDF.
        
        Args:
            data: Dictionary containing solution data
            output_path: Path to save PDF
        """
        pdf = FPDF()
        pdf.add_page()
        
        # Title
        pdf.set_font('Arial', 'B', 16)
        pdf.cell(0, 10, 'RecursiveLearn - Solution Report', ln=True, align='C')
        pdf.ln(5)
        
        # Relation
        pdf.set_font('Arial', 'B', 12)
        pdf.cell(0, 8, 'Recurrence Relation:', ln=True)
        pdf.set_font('Arial', '', 11)
        pdf.multi_cell(0, 6, data.get('relation', 'N/A'))
        pdf.ln(3)
        
        # Initial Conditions
        pdf.set_font('Arial', 'B', 12)
        pdf.cell(0, 8, 'Initial Conditions:', ln=True)
        pdf.set_font('Arial', '', 11)
        initial = data.get('initial_conditions', {})
        for key, val in initial.items():
            pdf.cell(0, 6, f'  a_{key} = {val}', ln=True)
        pdf.ln(3)
        
        # Closed Form Solution
        if 'closed_form' in data:
            pdf.set_font('Arial', 'B', 12)
            pdf.cell(0, 8, 'Closed-Form Solution:', ln=True)
            pdf.set_font('Arial', '', 11)
            pdf.multi_cell(0, 6, data['closed_form'])
            pdf.ln(3)
        
        # Sequence Terms
        if 'sequence' in data:
            pdf.set_font('Arial', 'B', 12)
            pdf.cell(0, 8, 'First 20 Terms:', ln=True)
            pdf.set_font('Arial', '', 11)
            seq_str = ', '.join(str(x) for x in data['sequence'][:20])
            pdf.multi_cell(0, 6, seq_str)
            pdf.ln(3)
        
        # Solution Steps
        if 'steps' in data:
            pdf.set_font('Arial', 'B', 12)
            pdf.cell(0, 8, 'Solution Steps:', ln=True)
            pdf.set_font('Arial', '', 10)
            
            for step in data['steps']:
                # Handle long lines
                if len(step) > 80:
                    pdf.multi_cell(0, 5, step)
                else:
                    pdf.cell(0, 5, step, ln=True)
        
        # Save PDF
        pdf.output(output_path)
        return output_path


class DOCXExporter:
    """Export solutions to DOCX format."""
    
    def export_solution(self, data: Dict, output_path: str):
        """
        Export a solution to DOCX.
        
        Args:
            data: Dictionary containing solution data
            output_path: Path to save DOCX
        """
        doc = Document()
        
        # Title
        title = doc.add_heading('RecursiveLearn - Solution Report', 0)
        title.alignment = 1  # Center
        
        # Relation
        doc.add_heading('Recurrence Relation', 2)
        doc.add_paragraph(data.get('relation', 'N/A'))
        
        # Initial Conditions
        doc.add_heading('Initial Conditions', 2)
        initial = data.get('initial_conditions', {})
        for key, val in initial.items():
            doc.add_paragraph(f'a_{key} = {val}', style='List Bullet')
        
        # Closed Form
        if 'closed_form' in data:
            doc.add_heading('Closed-Form Solution', 2)
            p = doc.add_paragraph(data['closed_form'])
            p.runs[0].font.bold = True
            p.runs[0].font.color.rgb = RGBColor(0, 0, 255)
        
        # Sequence
        if 'sequence' in data:
            doc.add_heading('First 20 Terms', 2)
            seq_str = ', '.join(str(x) for x in data['sequence'][:20])
            doc.add_paragraph(seq_str)
        
        # Steps
        if 'steps' in data:
            doc.add_heading('Solution Steps', 2)
            for step in data['steps']:
                doc.add_paragraph(step, style='List Number')
        
        # Save
        doc.save(output_path)
        return output_path


class GraphExporter:
    """Export graphs and visualizations."""
    
    @staticmethod
    def save_graph(fig, output_path: str):
        """Save matplotlib figure to file."""
        fig.savefig(output_path, dpi=300, bbox_inches='tight')
        return output_path
    
    @staticmethod
    def export_sequence_plot(sequence: List[int], title: str, output_path: str):
        """Create and save a sequence plot."""
        fig, ax = plt.subplots(figsize=(10, 6))
        
        n_values = list(range(len(sequence)))
        ax.plot(n_values, sequence, 'bo-', linewidth=2, markersize=8)
        
        ax.set_xlabel('n', fontsize=12)
        ax.set_ylabel('a_n', fontsize=12)
        ax.set_title(title, fontsize=14, fontweight='bold')
        ax.grid(True, alpha=0.3)
        
        fig.tight_layout()
        fig.savefig(output_path, dpi=300, bbox_inches='tight')
        plt.close(fig)
        
        return output_path
